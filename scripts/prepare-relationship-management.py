"""Read CEO DOCX/XLSX into a private RM bundle and COMPLETE SQL75 installation.
Run with the bundled workspace Python (python-docx + openpyxl). No network or DB writes.
Never commit generated output or print the bank, keys or behavior anchors.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

def prepare(folder, output):
    docpath = folder / 'BinaInsights_Laporan_Relationship_Management_v4.docx'
    bankpath = folder / 'BinaInsights_Relationship_Management_Bank_Soal_V4.xlsx'
    book = load_workbook(bankpath, read_only=True, data_only=False)
    def rows(sheet):
        data = [list(r) for r in book[sheet].iter_rows(values_only=True) if any(v is not None for v in r)]
        return [dict(zip(data[0], r)) for r in data[1:]]
    report = Document(docpath)
    tables = [t for t in report.tables if any('Kandidat layanan BinaHub' in c.text for c in t.rows[0].cells)]
    if len(tables) != 1: raise ValueError('RM_REPORT_MAPPING_TABLE_REQUIRED')
    mapping = {r.cells[0].text: [c.text for c in r.cells] for r in tables[0].rows[1:]}
    competencies = []
    for c in rows('Competency_Map'):
        name = c['Competency']
        info = mapping[name]
        candidates = []
        for line in info[2].splitlines():
            match = re.match(r'(SS-\d{2})\s*•\s*(.*?)\s*•', line)
            if not match: raise ValueError('RM_REPORT_CATALOG_MAPPING_INVALID')
            candidates.append({'code': match[1], 'name': match[2], 'alignment':line.split('•',2)[2].strip()})
        competencies.append({'code': c['Competency_ID'], 'name': name, 'definition': c['Working_Definition'],
            'behaviorAnchors': info[1].splitlines(), 'candidates': candidates, 'deliveryNote': info[3], 'validation': info[4]})
    byname = {c['name']: c['code'] for c in competencies}
    sr = rows('Self_Report_Items')
    scale = [{'value': n, 'label': sr[0]['Scale_'+str(n)]} for n in range(1, 6)]
    selfitems = []
    for i in sr:
        if i['Reverse_Scored'] != 'No' or any(i['Scale_'+str(n)] != scale[n-1]['label'] for n in range(1, 6)):
            raise ValueError('RM_SCALE_REVIEW_REQUIRED')
        selfitems.append({'id': i['Item_ID'], 'competencyCode': byname[i['Competency']], 'prompt': i['Item_Text'], 'behaviorAnchor': i['Key_Action_Focus']})
    sjt = []
    for i in rows('SJT_Items'):
        scores = {o: int(i['Score_'+o]) for o in 'ABCD'}
        if any(not isinstance(i['Score_'+o], (int,float)) or i['Score_'+o] != scores[o] or not 0 <= scores[o] <= 3 for o in 'ABCD'):
            raise ValueError('RM_CREDITS_INVALID')
        if scores[i['Best_Option']] != 3: raise ValueError('RM_KEY_INVALID')
        sjt.append({'id': i['Item_ID'], 'title': i['Scenario_Title'], 'scenario': i['Scenario_Text'],
            'options': [{'id': o, 'text': i['Option_'+o]} for o in 'ABCD'],
            'tags': [byname[t.strip()] for t in i['Competency_Tags'].split(';')],
            'scores': scores, 'bestOption': i['Best_Option'], 'rationale': i['Rationale']})
    if len(competencies) != 9 or len(selfitems) != 27 or len(sjt) != 13: raise ValueError('RM_COUNTS_INVALID')
    if any(sum(i['competencyCode'] == c['code'] for i in selfitems) != 3 for c in competencies): raise ValueError('RM_SR_COVERAGE_INVALID')
    if len({i['id'] for i in selfitems+sjt}) != 40: raise ValueError('RM_IDS_INVALID')
    version = '4.0.0-draft.1'
    bundle = {'metadata': {'version': version, 'status': 'draft', 'publicEnabled': False,
        'sources': [{'file': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in [docpath,bankpath]],
        'reportSelfReportCount':24,'bankSelfReportCount':27}, 'competencies':competencies,'scale':scale,
        'selfReportItems':selfitems,'sjtItems':sjt,'scoringPolicy':None}
    canonical = json.dumps(bundle, ensure_ascii=False, sort_keys=True, separators=(',',':'))
    digest = hashlib.sha256(canonical.encode()).hexdigest()
    output.mkdir(parents=True,exist_ok=True)
    jsonpath = output / 'relationship-management-v4.private.json'
    sqlpath = output / 'SQL75_RM_COMPLETE.private.sql'
    if jsonpath.exists() and jsonpath.read_text(encoding='utf-8') != canonical: raise ValueError('RM_EXISTING_OUTPUT_DIFFERS_USE_NEW_DIRECTORY')
    jsonpath.write_text(canonical,encoding='utf-8')
    schema = (Path(__file__).resolve().parents[1] / 'supabase/migrations/0075_relationship_management_foundation.sql').read_text(encoding='utf-8')
    schema = schema.removesuffix('commit;\n')
    escaped = canonical.replace("'","''")
    sql = '-- PRIVATE FULL INSTALLATION: schema + CEO source. Do not share or commit.\n'+schema+f"\nDO $$ BEGIN\n IF EXISTS(SELECT 1 FROM public.rm_instrument_versions WHERE version='{version}' AND bundle_sha256<>'{digest}') THEN RAISE EXCEPTION 'RM_IMMUTABLE_VERSION_CONFLICT'; END IF;\nEND; $$;\nINSERT INTO public.rm_instrument_versions(version,manifest,bundle_sha256) VALUES ('{version}','{escaped}'::jsonb,'{digest}') ON CONFLICT(version) DO NOTHING;\nCOMMIT;\nSELECT version, jsonb_array_length(manifest->'selfReportItems') AS self_report_count, jsonb_array_length(manifest->'sjtItems') AS sjt_count FROM public.rm_instrument_versions WHERE version='{version}';\n"
    sqlpath.write_text(sql,encoding='utf-8')
    print(json.dumps({'version':version,'selfReport':27,'sjt':13,'competencies':9,'publicEnabled':False,'sha256':digest,'privateSQL':str(sqlpath)},ensure_ascii=False))

if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--source-dir',type=Path,required=True)
    parser.add_argument('--output-dir',type=Path,required=True)
    args=parser.parse_args()
    prepare(args.source_dir,args.output_dir)
